#!/usr/bin/env python3
"""Export a canonical survey to human-readable, web, document, and interchange formats."""

from __future__ import annotations

import argparse
import csv
import html
import json
import sys
import textwrap
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import yaml

from generate_codebook import write_codebook
from survey_io import iter_question_contexts, load_survey, validate_survey

TEXTUAL_FORMATS = {"markdown", "text", "html", "xml", "json", "yaml"}


def _question_number_map(data: Mapping[str, Any]) -> dict[str, int]:
    return {str(context.question["id"]): index for index, context in enumerate(iter_question_contexts(data), 1)}


def _required_label(question: Mapping[str, Any]) -> str:
    return "Required" if question.get("required") else "Optional"


def _choice_lines(question: Mapping[str, Any], marker: str = "- ") -> list[str]:
    lines: list[str] = []
    choices = question.get("choices", [])
    if isinstance(choices, list):
        for choice in choices:
            if isinstance(choice, Mapping):
                suffix = " (exclusive)" if choice.get("exclusive") else ""
                lines.append(f"{marker}{choice.get('label', '')} [{choice.get('value', '')}]{suffix}")
    return lines


def export_markdown(data: Mapping[str, Any], *, include_notes: bool = True) -> str:
    survey = data["survey"]
    numbers = _question_number_map(data)
    lines = [
        f"# {survey['title']}",
        "",
        f"**Survey ID:** `{survey['id']}`  ",
        f"**Version:** {survey['version']}  ",
        f"**Audience:** {survey['target_audience']}  ",
        f"**Estimated time:** {survey['estimated_minutes']} minutes",
        "",
        f"> **Purpose:** {survey['objective']}",
        "",
        "## Before you begin",
        "",
        survey["consent"]["text"],
        "",
    ]
    for block in data["blocks"]:
        lines.extend([f"## {block['title']}", ""])
        if block.get("description"):
            lines.extend([str(block["description"]), ""])
        for question in block["questions"]:
            qid = str(question["id"])
            qtype = str(question["type"])
            lines.extend(
                [
                    f"### {numbers[qid]}. {question['text']}",
                    "",
                    f"`ID: {qid}` · `{qtype}` · {_required_label(question)}",
                    "",
                ]
            )
            if question.get("help_text"):
                lines.extend([f"_{question['help_text']}_", ""])
            if qtype in {"single_choice", "dropdown", "yes_no"}:
                lines.extend(f"- ○ {line[2:]}" for line in _choice_lines(question))
            elif qtype == "multi_choice":
                lines.extend(f"- ☐ {line[2:]}" for line in _choice_lines(question))
            elif qtype in {"likert", "rating", "nps", "slider", "number_scale"}:
                scale = question.get("scale", {})
                anchors = scale.get("anchors", {}) if isinstance(scale, Mapping) else {}
                lines.append(f"Scale: {scale.get('min')} to {scale.get('max')} in steps of {scale.get('step')}")
                if anchors:
                    lines.append("Anchors: " + "; ".join(f"{value} = {label}" for value, label in anchors.items()))
            elif qtype in {"matrix", "likert_matrix", "bipolar_matrix", "rank_order_matrix", "constant_sum_matrix", "text_entry_matrix", "side_by_side"}:
                rows = question.get("rows", [])
                columns = question.get("columns", [])
                lines.append("Rows: " + "; ".join(str(row.get("label", "")) for row in rows if isinstance(row, Mapping)))
                lines.append("Columns: " + "; ".join(str(column.get("label", "")) for column in columns if isinstance(column, Mapping)))
            elif qtype in {"text", "numeric", "date", "calendar"}:
                lines.append("`[Short response]`")
            elif qtype == "textarea":
                lines.append("`[Long response]`")
            elif qtype in {"rank", "constant_sum", "maxdiff", "pick_group_rank", "drilldown"}:
                lines.extend(_choice_lines(question))
            elif qtype == "descriptive":
                pass
            else:
                lines.append(f"`[{qtype} response]`")
            if include_notes:
                lines.extend(
                    [
                        "",
                        f"**Rationale:** {question['rationale']}",
                        f"**Construct:** {question['construct']}",
                        f"**Analysis use:** {question['analysis_use']}",
                    ]
                )
                if question.get("display_if"):
                    lines.append(f"**Display if:** `{json.dumps(question['display_if'], ensure_ascii=False)}`")
                if question.get("skip_to"):
                    lines.append(f"**Skip to:** `{question['skip_to']}`")
            lines.append("")
    lines.extend(["## Thank you", "", "Thank you for sharing your experience.", ""])
    return "\n".join(lines)


def export_text(data: Mapping[str, Any], *, include_notes: bool = True) -> str:
    survey = data["survey"]
    numbers = _question_number_map(data)
    lines = [
        survey["title"].upper(),
        "=" * len(survey["title"]),
        f"Survey ID: {survey['id']}",
        f"Version: {survey['version']}",
        f"Audience: {survey['target_audience']}",
        f"Estimated time: {survey['estimated_minutes']} minutes",
        "",
        f"Purpose: {survey['objective']}",
        "",
        "BEFORE YOU BEGIN",
        survey["consent"]["text"],
        "",
    ]
    for block in data["blocks"]:
        lines.extend([block["title"].upper(), "-" * len(block["title"])])
        if block.get("description"):
            lines.append(str(block["description"]))
        lines.append("")
        for question in block["questions"]:
            qid = str(question["id"])
            lines.append(f"{numbers[qid]}. {question['text']}")
            lines.append(f"   ID={qid} TYPE={question['type']} {_required_label(question).upper()}")
            if question.get("help_text"):
                lines.append(f"   Help: {question['help_text']}")
            qtype = question["type"]
            if qtype in {"single_choice", "dropdown", "yes_no", "multi_choice", "rank", "constant_sum", "maxdiff", "pick_group_rank", "drilldown"}:
                for choice in question.get("choices", []):
                    box = "[ ]" if qtype == "multi_choice" else "( )"
                    lines.append(f"   {box} {choice['label']} [{choice['value']}]")
            elif qtype in {"likert", "rating", "nps", "slider", "number_scale"}:
                scale = question.get("scale", {})
                lines.append(f"   Scale {scale.get('min')}–{scale.get('max')}, step {scale.get('step')}")
                for value, label in scale.get("anchors", {}).items():
                    lines.append(f"   {value}: {label}")
            elif qtype in {"matrix", "likert_matrix", "bipolar_matrix", "rank_order_matrix", "constant_sum_matrix", "text_entry_matrix", "side_by_side"}:
                lines.append("   Rows: " + "; ".join(str(row["label"]) for row in question.get("rows", [])))
                lines.append("   Columns: " + "; ".join(str(column["label"]) for column in question.get("columns", [])))
            elif qtype == "textarea":
                lines.append("   [Long response]")
            elif qtype != "descriptive":
                lines.append("   [Response]")
            if include_notes:
                lines.append(f"   Rationale: {question['rationale']}")
                lines.append(f"   Analysis: {question['analysis_use']}")
            lines.append("")
    lines.append("THANK YOU")
    return "\n".join(lines) + "\n"


def _html_choice_inputs(question: Mapping[str, Any], input_type: str) -> str:
    qid = html.escape(str(question["id"]), quote=True)
    required = " required" if question.get("required") and input_type == "radio" else ""
    parts: list[str] = []
    for index, choice in enumerate(question.get("choices", []), 1):
        control_id = f"{qid}_{index}"
        value = html.escape(str(choice.get("value", "")), quote=True)
        label = html.escape(str(choice.get("label", "")))
        exclusive = " data-exclusive=\"true\"" if choice.get("exclusive") else ""
        parts.append(
            f'<div class="option"><input id="{control_id}" name="{qid}" type="{input_type}" value="{value}"{required}{exclusive}>'
            f'<label for="{control_id}">{label}</label></div>'
        )
    return "\n".join(parts)


def _html_scale(question: Mapping[str, Any]) -> str:
    qid = html.escape(str(question["id"]), quote=True)
    scale = question.get("scale", {})
    minimum = int(scale.get("min", 1))
    maximum = int(scale.get("max", 5))
    step = int(scale.get("step", 1))
    anchors = {str(k): str(v) for k, v in scale.get("anchors", {}).items()}
    required = " required" if question.get("required") else ""
    parts = ['<div class="scale" role="radiogroup">']
    for value in range(minimum, maximum + 1, step):
        control_id = f"{qid}_{value}"
        label_text = anchors.get(str(value), "")
        parts.append(
            f'<div class="scale-point"><input id="{control_id}" name="{qid}" type="radio" value="{value}"{required}>'
            f'<label for="{control_id}"><span class="value">{value}</span><span>{html.escape(label_text)}</span></label></div>'
        )
    parts.append("</div>")
    return "\n".join(parts)


def _html_matrix(question: Mapping[str, Any]) -> str:
    qid = html.escape(str(question["id"]), quote=True)
    rows = [row for row in question.get("rows", []) if isinstance(row, Mapping)]
    columns = [column for column in question.get("columns", []) if isinstance(column, Mapping)]
    multiple = question.get("type") in {"constant_sum_matrix", "text_entry_matrix"}
    input_type = "text" if question.get("type") == "text_entry_matrix" else ("checkbox" if multiple else "radio")
    parts = ['<div class="table-wrap"><table>', f'<caption class="sr-only">{html.escape(str(question["text"]))}</caption>', "<thead><tr><th scope=\"col\">Item</th>"]
    for column in columns:
        parts.append(f'<th scope="col">{html.escape(str(column.get("label", "")))}</th>')
    parts.append("</tr></thead><tbody>")
    for r_index, row in enumerate(rows, 1):
        parts.append(f'<tr><th scope="row">{html.escape(str(row.get("label", "")))}</th>')
        for c_index, column in enumerate(columns, 1):
            control_id = f"{qid}_r{r_index}_c{c_index}"
            name = f"{qid}_{row.get('id', r_index)}"
            value = html.escape(str(column.get("value", "")), quote=True)
            label = f"{row.get('label', '')}: {column.get('label', '')}"
            required = " required" if question.get("required") and input_type == "radio" and c_index == 1 else ""
            parts.append(
                f'<td><label class="sr-only" for="{control_id}">{html.escape(label)}</label>'
                f'<input id="{control_id}" name="{html.escape(name, quote=True)}" type="{input_type}" value="{value}"{required}></td>'
            )
        parts.append("</tr>")
    parts.append("</tbody></table></div>")
    return "\n".join(parts)


def export_html(data: Mapping[str, Any]) -> str:
    survey = data["survey"]
    title = html.escape(str(survey["title"]))
    parts = [
        "<!doctype html>",
        '<html lang="en">',
        "<head>",
        '<meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        f"<title>{title}</title>",
        "<style>",
        ":root{font-family:system-ui,-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;line-height:1.5;color:#171717;background:#f5f5f5}",
        "body{margin:0}.shell{max-width:760px;margin:0 auto;padding:2rem 1rem 5rem}header,section,.consent{background:#fff;border:1px solid #ddd;border-radius:.75rem;padding:1.25rem;margin-bottom:1rem}",
        "fieldset{border:0;padding:0;margin:0}legend{font-weight:700;font-size:1.05rem;margin-bottom:.5rem}.meta{color:#555;font-size:.9rem}.required{font-weight:600}.option{display:flex;gap:.65rem;align-items:flex-start;padding:.35rem 0}",
        "input,select,textarea{font:inherit}input[type=text],input[type=number],input[type=date],select,textarea{width:100%;box-sizing:border-box;padding:.65rem;border:1px solid #777;border-radius:.35rem}textarea{min-height:8rem}",
        ".question+.question{margin-top:1.75rem}.scale{display:grid;grid-template-columns:repeat(auto-fit,minmax(5rem,1fr));gap:.5rem}.scale-point{display:flex;flex-direction:column;gap:.3rem}.scale-point label{display:flex;flex:1;flex-direction:column;align-items:center;text-align:center;border:1px solid #aaa;border-radius:.35rem;padding:.4rem;overflow-wrap:anywhere}.value{font-weight:700}button{font:inherit;padding:.7rem 1rem;border:1px solid #171717;border-radius:.35rem;background:#171717;color:#fff}input:focus-visible,textarea:focus-visible,select:focus-visible,button:focus-visible{outline:3px solid #2458a6;outline-offset:3px}",
        ".table-wrap{overflow-x:auto}table{border-collapse:collapse;width:100%}th,td{border:1px solid #bbb;padding:.55rem;text-align:center}th[scope=row]{text-align:left}.sr-only{position:absolute;width:1px;height:1px;padding:0;margin:-1px;overflow:hidden;clip:rect(0,0,0,0);white-space:nowrap;border:0}",
        "@media print{body{background:#fff}.shell{max-width:none;padding:0}header,section,.consent{border:0;break-inside:avoid}.meta{color:#000}}",
        "</style>",
        "</head>",
        "<body><main class=\"shell\">",
        f"<header><h1>{title}</h1><p>{html.escape(str(survey['objective']))}</p><p class=\"meta\">Estimated time: {survey['estimated_minutes']} minutes · Survey ID: <code>{html.escape(str(survey['id']))}</code></p></header>",
        f"<div class=\"consent\"><h2>Before you begin</h2><p>{html.escape(str(survey['consent']['text']))}</p><p class=\"meta\">Questions: {html.escape(str(survey['consent']['contact']))}</p></div>",
        '<section aria-labelledby="preview-title"><h2 id="preview-title">Review preview</h2><p>Try the controls with fictional answers. This file does not collect, save or send responses. Branching and platform-specific rules are not applied here.</p></section>',
        '<form aria-label="Survey preview" method="dialog">',
    ]
    for block in data["blocks"]:
        parts.append(f'<section id="block-{html.escape(str(block["id"]), quote=True)}"><h2>{html.escape(str(block["title"]))}</h2>')
        if block.get("description"):
            parts.append(f"<p>{html.escape(str(block['description']))}</p>")
        for question in block["questions"]:
            qid = html.escape(str(question["id"]), quote=True)
            help_id = f"{qid}_help"
            required_text = '<span class="required">Required</span>' if question.get("required") else "Optional"
            qtype = str(question["type"])
            parts.append(f'<div class="question" data-question-id="{qid}">')
            if qtype == "descriptive":
                parts.append(f"<div>{html.escape(str(question['text']))}</div>")
                parts.append("</div>")
                continue
            parts.append(f'<fieldset><legend>{html.escape(str(question["text"]))}</legend><div class="meta">{required_text} · ID: <code>{qid}</code></div>')
            if question.get("help_text"):
                parts.append(f'<p id="{help_id}" class="meta">{html.escape(str(question["help_text"]))}</p>')
            if qtype in {"single_choice", "yes_no"}:
                parts.append(_html_choice_inputs(question, "radio"))
            elif qtype == "multi_choice":
                parts.append(_html_choice_inputs(question, "checkbox"))
            elif qtype == "dropdown":
                required = " required" if question.get("required") else ""
                parts.append(f'<label class="sr-only" for="{qid}_select">{html.escape(str(question["text"]))}</label><select id="{qid}_select" name="{qid}"{required}><option value="">Select one</option>')
                for choice in question.get("choices", []):
                    parts.append(f'<option value="{html.escape(str(choice.get("value", "")), quote=True)}">{html.escape(str(choice.get("label", "")))}</option>')
                parts.append("</select>")
            elif qtype in {"likert", "rating", "nps", "number_scale"}:
                parts.append(_html_scale(question))
            elif qtype == "slider":
                scale = question["scale"]
                required = " required" if question.get("required") else ""
                parts.append(f'<label for="{qid}_slider">{html.escape(str(question["text"]))}</label><input id="{qid}_slider" name="{qid}" type="range" min="{scale["min"]}" max="{scale["max"]}" step="{scale["step"]}"{required}><output for="{qid}_slider">{scale["min"]}</output>')
            elif qtype in {"matrix", "likert_matrix", "bipolar_matrix", "rank_order_matrix", "constant_sum_matrix", "text_entry_matrix", "side_by_side"}:
                parts.append(_html_matrix(question))
            elif qtype in {"text", "numeric", "date", "calendar"}:
                input_type = {"numeric": "number", "date": "date", "calendar": "datetime-local"}.get(qtype, "text")
                required = " required" if question.get("required") else ""
                parts.append(f'<label class="sr-only" for="{qid}_input">{html.escape(str(question["text"]))}</label><input id="{qid}_input" name="{qid}" type="{input_type}"{required}>')
            elif qtype == "textarea":
                required = " required" if question.get("required") else ""
                parts.append(f'<label class="sr-only" for="{qid}_input">{html.escape(str(question["text"]))}</label><textarea id="{qid}_input" name="{qid}"{required}></textarea>')
            elif qtype in {"rank", "constant_sum", "maxdiff", "pick_group_rank", "drilldown"}:
                parts.append(_html_choice_inputs(question, "checkbox" if qtype in {"pick_group_rank", "maxdiff"} else "text"))
            else:
                parts.append(f'<p class="meta">Interactive {html.escape(qtype)} control must be programmed in the destination platform. Provide an accessible fallback.</p>')
            parts.append("</fieldset></div>")
        parts.append("</section>")
    parts.extend([
        '<section><h2>Check the preview</h2><button type="button" id="check-preview" hidden>Check required answers</button><p id="preview-status" role="status" aria-live="polite">Answers stay in this page only. This is not a submission.</p><noscript>JavaScript is required for the local answer check. No responses are collected.</noscript></section>',
        "</form>",
        '<script>const form=document.querySelector("form");const status=document.getElementById("preview-status");const check=document.getElementById("check-preview");check.hidden=false;form.addEventListener("submit",event=>event.preventDefault());form.addEventListener("input",()=>{status.textContent="Answers changed. Check again when ready. Nothing has been sent or saved.";});check.addEventListener("click",()=>{const valid=form.reportValidity();status.textContent=valid?"Required answers are present. Nothing has been sent or saved. Branching and platform rules still need destination testing.":"Complete the required answers highlighted above. Nothing has been sent or saved.";});</script>',
        "</main></body></html>",
    ])
    return "\n".join(parts) + "\n"


def export_xml(data: Mapping[str, Any]) -> str:
    root = ET.Element("survey", {"id": str(data["survey"]["id"]), "version": str(data["survey"]["version"])})
    metadata = ET.SubElement(root, "metadata")
    for key in ("title", "objective", "target_audience", "mode", "estimated_minutes"):
        node = ET.SubElement(metadata, key)
        node.text = str(data["survey"].get(key, ""))
    analysis = ET.SubElement(root, "analysisPlan")
    for field in ("decisions", "primary_metrics", "secondary_metrics", "segments", "exclusions"):
        container = ET.SubElement(analysis, field)
        for value in data.get("analysis_plan", {}).get(field, []) or []:
            item = ET.SubElement(container, "item")
            item.text = str(value)
    blocks = ET.SubElement(root, "blocks")
    for block in data["blocks"]:
        bnode = ET.SubElement(blocks, "block", {"id": str(block["id"]), "title": str(block["title"])})
        for question in block["questions"]:
            attrs = {
                "id": str(question["id"]),
                "exportTag": str(question.get("export_tag", question["id"])),
                "type": str(question["type"]),
                "required": str(bool(question.get("required"))).lower(),
            }
            qnode = ET.SubElement(bnode, "question", attrs)
            for key in ("text", "help_text", "rationale", "construct", "analysis_use"):
                if question.get(key) is not None:
                    child = ET.SubElement(qnode, key)
                    child.text = str(question.get(key, ""))
            for field_name, tag in (("choices", "choice"), ("rows", "row"), ("columns", "column")):
                if question.get(field_name):
                    container = ET.SubElement(qnode, field_name)
                    for option in question[field_name]:
                        node = ET.SubElement(container, tag, {"id": str(option["id"]), "value": str(option["value"])})
                        node.text = str(option["label"])
            if question.get("scale"):
                scale = question["scale"]
                snode = ET.SubElement(qnode, "scale", {"min": str(scale["min"]), "max": str(scale["max"]), "step": str(scale["step"])})
                for value, label in scale.get("anchors", {}).items():
                    anchor = ET.SubElement(snode, "anchor", {"value": str(value)})
                    anchor.text = str(label)
            if question.get("display_if"):
                logic = ET.SubElement(qnode, "displayIf")
                logic.text = json.dumps(question["display_if"], ensure_ascii=False, sort_keys=True)
            if question.get("skip_to"):
                skip = ET.SubElement(qnode, "skipTo")
                skip.text = str(question["skip_to"])
    try:
        ET.indent(root, space="  ")
    except AttributeError:  # pragma: no cover
        pass
    return '<?xml version="1.0" encoding="UTF-8"?>\n' + ET.tostring(root, encoding="unicode") + "\n"


def export_docx(data: Mapping[str, Any], output: Path, *, include_notes: bool = True) -> None:
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.text import WD_BREAK
        from docx.shared import Inches, Pt
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("DOCX export requires python-docx: python -m pip install python-docx") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    title = document.add_heading(str(data["survey"]["title"]), level=0)
    document.add_paragraph(f"Survey ID: {data['survey']['id']}  |  Version: {data['survey']['version']}")
    document.add_paragraph(f"Audience: {data['survey']['target_audience']}")
    document.add_paragraph(f"Estimated time: {data['survey']['estimated_minutes']} minutes")
    purpose = document.add_paragraph()
    purpose.add_run("Purpose: ").bold = True
    purpose.add_run(str(data["survey"]["objective"]))
    document.add_heading("Before you begin", level=1)
    document.add_paragraph(str(data["survey"]["consent"]["text"]))
    numbers = _question_number_map(data)
    for block_index, block in enumerate(data["blocks"]):
        document.add_heading(str(block["title"]), level=1)
        if block.get("description"):
            document.add_paragraph(str(block["description"]))
        for question in block["questions"]:
            qid = str(question["id"])
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{numbers[qid]}. {question['text']}").bold = True
            meta = document.add_paragraph(f"ID: {qid} | Type: {question['type']} | {_required_label(question)}")
            meta.style = document.styles["Caption"]
            if question.get("help_text"):
                document.add_paragraph(str(question["help_text"]), style="Intense Quote")
            qtype = question["type"]
            if qtype in {"single_choice", "dropdown", "yes_no", "multi_choice", "rank", "constant_sum", "maxdiff", "pick_group_rank", "drilldown"}:
                for choice in question.get("choices", []):
                    marker = "☐" if qtype == "multi_choice" else "○"
                    document.add_paragraph(f"{marker} {choice['label']} [{choice['value']}]", style="List Bullet")
            elif qtype in {"likert", "rating", "nps", "slider", "number_scale"}:
                scale = question["scale"]
                table = document.add_table(rows=2, cols=len(range(int(scale["min"]), int(scale["max"]) + 1, int(scale["step"]))))
                table.style = "Table Grid"
                for index, value in enumerate(range(int(scale["min"]), int(scale["max"]) + 1, int(scale["step"]))):
                    table.cell(0, index).text = str(value)
                    table.cell(1, index).text = str(scale.get("anchors", {}).get(str(value), ""))
            elif qtype in {"matrix", "likert_matrix", "bipolar_matrix", "rank_order_matrix", "constant_sum_matrix", "text_entry_matrix", "side_by_side"}:
                rows = question.get("rows", [])
                columns = question.get("columns", [])
                table = document.add_table(rows=len(rows) + 1, cols=len(columns) + 1)
                table.style = "Table Grid"
                for c_index, column in enumerate(columns, 1):
                    table.cell(0, c_index).text = str(column["label"])
                for r_index, row in enumerate(rows, 1):
                    table.cell(r_index, 0).text = str(row["label"])
                    for c_index in range(1, len(columns) + 1):
                        table.cell(r_index, c_index).text = "○"
            elif qtype == "textarea":
                document.add_paragraph("\n" + "_" * 80 + "\n" + "_" * 80)
            elif qtype != "descriptive":
                document.add_paragraph("_" * 80)
            if include_notes:
                note = document.add_paragraph()
                note.add_run("Rationale: ").bold = True
                note.add_run(str(question["rationale"]))
                note.add_run("\nAnalysis use: ").bold = True
                note.add_run(str(question["analysis_use"]))
        if block_index < len(data["blocks"]) - 1:
            document.add_page_break()
    document.add_heading("Thank you", level=1)
    document.add_paragraph("Thank you for sharing your experience.")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def export_pdf(data: Mapping[str, Any], output: Path, *, include_notes: bool = True) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        from xml.sax.saxutils import escape
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("PDF export requires reportlab: python -m pip install reportlab") from exc

    output.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(output),
        pagesize=LETTER,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
        title=str(data["survey"]["title"]),
        author="Survey Strategist",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Meta", parent=styles["Normal"], fontSize=8.5, leading=11, textColor=colors.HexColor("#444444")))
    styles.add(ParagraphStyle(name="Question", parent=styles["Heading3"], fontSize=10.5, leading=14, spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle(name="Scale", parent=styles["Normal"], fontSize=8, leading=10, alignment=TA_CENTER))
    story: list[Any] = []
    story.append(Paragraph(escape(str(data["survey"]["title"])), styles["Title"]))
    story.append(Paragraph(escape(f"Survey ID: {data['survey']['id']} | Version: {data['survey']['version']} | Estimated time: {data['survey']['estimated_minutes']} minutes"), styles["Meta"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(f"<b>Purpose:</b> {escape(str(data['survey']['objective']))}", styles["Normal"]))
    story.append(Paragraph(f"<b>Audience:</b> {escape(str(data['survey']['target_audience']))}", styles["Normal"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Before you begin", styles["Heading2"]))
    story.append(Paragraph(escape(str(data["survey"]["consent"]["text"])), styles["Normal"]))
    numbers = _question_number_map(data)
    for block_index, block in enumerate(data["blocks"]):
        story.append(Paragraph(escape(str(block["title"])), styles["Heading2"]))
        if block.get("description"):
            story.append(Paragraph(escape(str(block["description"])), styles["Normal"]))
        for question in block["questions"]:
            qid = str(question["id"])
            story.append(Paragraph(f"{numbers[qid]}. {escape(str(question['text']))}", styles["Question"]))
            story.append(Paragraph(escape(f"ID: {qid} | Type: {question['type']} | {_required_label(question)}"), styles["Meta"]))
            if question.get("help_text"):
                story.append(Paragraph(f"<i>{escape(str(question['help_text']))}</i>", styles["Normal"]))
            qtype = question["type"]
            if qtype in {"single_choice", "dropdown", "yes_no", "multi_choice", "rank", "constant_sum", "maxdiff", "pick_group_rank", "drilldown"}:
                for choice in question.get("choices", []):
                    marker = "[ ]" if qtype == "multi_choice" else "( )"
                    story.append(Paragraph(escape(f"{marker} {choice['label']} [{choice['value']}]"), styles["Normal"]))
            elif qtype in {"likert", "rating", "nps", "slider", "number_scale"}:
                scale = question["scale"]
                values = list(range(int(scale["min"]), int(scale["max"]) + 1, int(scale["step"])))
                table_data = [[str(value) for value in values], [Paragraph(escape(str(scale.get("anchors", {}).get(str(value), ""))), styles["Scale"]) for value in values]]
                table = Table(table_data, colWidths=[doc.width / len(values)] * len(values), repeatRows=0)
                table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 7), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
                story.append(table)
            elif qtype in {"matrix", "likert_matrix", "bipolar_matrix", "rank_order_matrix", "constant_sum_matrix", "text_entry_matrix", "side_by_side"}:
                columns = question.get("columns", [])
                table_data = [["Item"] + [str(column["label"]) for column in columns]]
                for row in question.get("rows", []):
                    table_data.append([str(row["label"])] + ["○" for _ in columns])
                table = Table(table_data, repeatRows=1)
                table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")), ("FONTSIZE", (0, 0), (-1, -1), 7), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
                story.append(table)
            elif qtype == "textarea":
                for _ in range(2):
                    story.append(Spacer(1, 16))
                    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.grey))
            elif qtype != "descriptive":
                story.append(Spacer(1, 16))
                story.append(HRFlowable(width="100%", thickness=0.5, color=colors.grey))
            if include_notes:
                story.append(Paragraph(f"<b>Rationale:</b> {escape(str(question['rationale']))}<br/><b>Analysis use:</b> {escape(str(question['analysis_use']))}", styles["Meta"]))
            story.append(Spacer(1, 4))
        if block_index < len(data["blocks"]) - 1:
            story.append(PageBreak())
    story.append(Paragraph("Thank you", styles["Heading2"]))
    story.append(Paragraph("Thank you for sharing your experience.", styles["Normal"]))
    def page_footer(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(document.leftMargin, 0.3 * inch, f"{data['survey']['id']} | Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=page_footer, onLaterPages=page_footer)


def export_to_path(data: Mapping[str, Any], output_format: str, output: str | Path, *, include_notes: bool = True) -> list[str]:
    """Export survey to a path and return explicit loss/compatibility warnings."""
    validate_survey(data)
    fmt = output_format.lower().replace(".md", "markdown")
    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    warnings: list[str] = []
    if fmt == "markdown":
        path.write_text(export_markdown(data, include_notes=include_notes), encoding="utf-8")
    elif fmt in {"text", "txt"}:
        path.write_text(export_text(data, include_notes=include_notes), encoding="utf-8")
    elif fmt == "html":
        path.write_text(export_html(data), encoding="utf-8")
    elif fmt == "xml":
        path.write_text(export_xml(data), encoding="utf-8")
    elif fmt == "json":
        path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    elif fmt in {"yaml", "yml"}:
        path.write_text(yaml.safe_dump(dict(data), sort_keys=False, allow_unicode=True), encoding="utf-8")
    elif fmt == "csv":
        write_codebook(data, path, "csv")
        warnings.append("CSV output is a codebook, not a runnable survey or response-data export.")
    elif fmt == "docx":
        export_docx(data, path, include_notes=include_notes)
    elif fmt == "pdf":
        export_pdf(data, path, include_notes=include_notes)
    elif fmt in {"qualtrics-advanced", "qualtrics-simple"}:
        try:
            from qualtrics_txt import export_advanced_txt, export_simple_txt
        except ImportError as exc:
            raise RuntimeError("Qualtrics export module is missing from this skill package.") from exc
        text, qwarnings = export_advanced_txt(data) if fmt == "qualtrics-advanced" else export_simple_txt(data)
        path.write_text(text, encoding="utf-8", newline="\n")
        warnings.extend(qwarnings)
    else:
        raise ValueError(f"Unsupported export format: {output_format}")

    manifest = path.with_suffix(".warnings.json")
    if warnings:
        manifest.write_text(
            json.dumps(
                {"format": fmt, "output": path.name, "warnings": warnings},
                indent=2,
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
        )
    elif manifest.exists():
        manifest.unlink()
    return warnings


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export a canonical survey to Markdown, text, accessible HTML, XML, JSON, YAML, CSV codebook, DOCX, PDF, or Qualtrics TXT.")
    parser.add_argument("survey", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--format", required=True, choices=("markdown", "text", "html", "xml", "json", "yaml", "csv", "docx", "pdf", "qualtrics-advanced", "qualtrics-simple"), dest="output_format")
    parser.add_argument("--respondent-only", action="store_true", help="Omit rationale and analysis notes from human-readable outputs")
    args = parser.parse_args(argv)
    try:
        data = load_survey(args.survey)
        warnings = export_to_path(data, args.output_format, args.output, include_notes=not args.respondent_only)
    except Exception as exc:
        print(f"Export failed: {exc}", file=sys.stderr)
        return 2
    print(args.output)
    for warning in warnings:
        print(f"WARNING: {warning}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
